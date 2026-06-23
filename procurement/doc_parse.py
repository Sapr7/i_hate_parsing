# -*- coding: utf-8 -*-
"""Extract line items from docx/xlsx attachments (table-first parsing)."""
from __future__ import annotations

import io
import re
import zipfile
from typing import Dict, List, Optional, Tuple

from procurement.doc_convert import convert_doc_to_docx, is_ole_doc
from procurement.http_utils import fetch_bytes

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

LINE_RE = re.compile(
    r"(?P<name>.{8,140}?)\s+(?:шт|штук|компл|упак)\.?\s+"
    r"(?P<qty>[\d\s,\.]+)\s+(?P<price>[\d\s,\.]{3,})",
    re.I,
)

GOOD_DOC = (
    "том 1",
    "том 2",
    "озп",
    "специфика",
    "ведомость",
    "перечень",
    "смет",
    "обоснование",
    "закупочная документация",
    "приложение",
    "форма",
)
BAD_DOC = (
    "проект договора",
    "извещение",
    "протокол",
    "аналитическ",
    "разъяснен",
    "разъясн",
    "преддоговор",
    "отчет о посещ",
)

HEADER_HINTS = {
    "name": ("наименован", "предмет", "товар", "описание", "характеристик", "материал", "номенклатур"),
    "qty": ("колич", "кол-во", "объем", "объём", "объем"),
    # price/total before unit — «Цена за ед.» must not match as unit
    "price": ("цена за", "цена, за", "цена ед", "стоимость за", "цена без", "цена с ндс", "цена руб", "цена"),
    "total": ("стоимость", "сумма", "всего", "общая цена"),
    "unit": ("ед. изм", "ед изм", "единиц измер", "единиц", "измерения"),
}

PRICE_HEADER_MARKERS = ("цена", "стоимость", "руб", "ндс", "сумма")
QTY_PRICE_MERGED_RE = re.compile(r"^(\d+)\s+(\d{3,}(?:[,\.]\d+)?)$")

BOILERPLATE = re.compile(
    r"договор|заказчик|исполнител|налог|сборов|актами|"
    r"настоящ|размещен|захоронен|ветеринар|"
    r"цена за ед\.?\s*изм|общая цена.*руб.*ндс",
    re.I,
)


def pick_documents(docs: List[Dict], limit: int = 4) -> List[Dict]:
    scored: List[Tuple[int, Dict]] = []
    for d in docs:
        title = (d.get("title") or "").lower()
        kind = (d.get("kind") or "").lower()
        if kind not in ("docx", "doc", "xlsx", "xls", "filestore"):
            continue
        score = 0
        for kw in GOOD_DOC:
            if kw in title:
                score += 3
        for kw in BAD_DOC:
            if kw in title:
                score -= 8
        if "проект договора" in title and "том 1" in title:
            score += 4
        if "том 2" in title and "технич" in title:
            score -= 3
        if "xlsx" in title or kind == "xlsx":
            score += 1
        scored.append((score, d))
    scored.sort(key=lambda x: (-x[0], x[1].get("title", "")))
    positive = [d for s, d in scored if s > 0]
    if positive:
        return positive[:limit]
    non_bad = [d for s, d in scored if s > -5][:limit]
    return non_bad or [d for _, d in scored[:limit]]


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _map_columns(header: List[str]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for i, cell in enumerate(header):
        low = _norm(cell)
        if not low:
            continue
        for field, hints in HEADER_HINTS.items():
            if field in mapping:
                continue
            if field == "unit" and any(x in low for x in PRICE_HEADER_MARKERS):
                continue
            if field == "unit" and "ед." in low and "изм" not in low and "единиц" not in low:
                continue
            if any(h in low for h in hints):
                mapping[field] = i
    return mapping


def _parse_price(value: str) -> Optional[float]:
    s = (value or "").strip()
    if not s:
        return None
    try:
        return float(s.replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _looks_like_unit(value: str) -> bool:
    value = (value or "").strip()
    if not value:
        return False
    low = value.lower()
    if re.search(r"шт|компл|упак|рул|рейс|усл|кг|м\b|м2|м3|л\b", low):
        return True
    if re.match(r"^[\d\s,\.]+$", value):
        return False
    parsed = _parse_price(value)
    if parsed is not None and parsed >= 50:
        return False
    return len(re.sub(r"\d", "", value)) >= 2


def _split_qty_price(qty: str, price: str) -> Tuple[str, str]:
    qty = (qty or "").strip()
    price = (price or "").strip()
    m = QTY_PRICE_MERGED_RE.match(qty)
    if not m:
        return qty, price
    q_val, p_val = m.group(1), m.group(2)
    p_norm = re.sub(r"\s+", "", price)
    q_p_norm = re.sub(r"\s+", "", p_val)
    q_norm = re.sub(r"\s+", "", qty)
    if not price or p_norm == q_p_norm or p_norm in q_norm:
        return q_val, price or p_val
    return qty, price


def _cell(row: List[str], idx: Optional[int]) -> str:
    if idx is None or idx >= len(row):
        return ""
    return row[idx].strip()


def _is_valid_position(name: str, qty: str, price: str) -> bool:
    name = (name or "").strip()
    qty = (qty or "").strip()
    price = (price or "").strip()
    if len(name) < 8:
        return False
    if len(name) > 220 or name.count("\n") > 1:
        return False
    if re.search(r"перечень документации|товарная накладная|торг-12|счет-фактура", name, re.I):
        return False
    if BOILERPLATE.search(name):
        return False
    if re.match(r"^\d{1,2}$", price):
        return False
    if re.match(r"^[\.,]+$", price):
        return False
    if len(re.sub(r"\D", "", qty)) > 12:
        return False
    if not qty and not price:
        return False
    if name.lower().startswith(("№", "n ", "п/п", "п.п")):
        return False
    if re.match(r"^[\d\s,\.]+$", name):
        return False
    if re.match(r"^(спецификация|итого|всего|примечание|наименование)$", name, re.I):
        return False
    return True


def _rows_from_table(matrix: List[List[str]], source: str) -> List[Dict[str, str]]:
    if len(matrix) < 2:
        return []
    header_idx = 0
    best_map: Dict[str, int] = {}
    for i, row in enumerate(matrix[:6]):
        m = _map_columns(row)
        if len(m) >= 2 and ("name" in m or "qty" in m or "price" in m):
            header_idx = i
            best_map = m
            break
    if not best_map:
        return []

    out: List[Dict[str, str]] = []
    for row in matrix[header_idx + 1 :]:
        if not any(c.strip() for c in row):
            continue
        name = _cell(row, best_map.get("name"))
        if not name and best_map.get("name") is None and len(row) > 1:
            name = _cell(row, 1 if _cell(row, 0).isdigit() else 0)
        qty = _cell(row, best_map.get("qty"))
        price = _cell(row, best_map.get("price"))
        unit = _cell(row, best_map.get("unit"))
        if unit and qty and _looks_like_unit(unit) and unit.lower() not in qty.lower():
            qty = f"{qty} {unit}".strip()
        qty, price = _split_qty_price(qty, price)
        if not _is_valid_position(name, qty, price):
            continue
        out.append(
            {
                "item": name,
                "quantity": qty,
                "price_for_one": price,
                "source": source,
            }
        )
    return out


def parse_docx_tables(raw: bytes, source: str) -> List[Dict[str, str]]:
    if Document is None:
        return []
    try:
        doc = Document(io.BytesIO(raw))
    except Exception:
        return []
    rows: List[Dict[str, str]] = []
    for table in doc.tables:
        matrix = [[c.text.strip() for c in tr.cells] for tr in table.rows]
        rows.extend(_rows_from_table(matrix, source))
    return rows


def parse_xlsx_tables(raw: bytes, source: str) -> List[Dict[str, str]]:
    if openpyxl is None:
        return []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return []
    rows: List[Dict[str, str]] = []
    for sheet in wb.worksheets:
        matrix: List[List[str]] = []
        for row in sheet.iter_rows(max_row=500, values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                matrix.append(cells)
        rows.extend(_rows_from_table(matrix, source))
    wb.close()
    return rows


def _regex_fallback(raw: bytes, source: str) -> List[Dict[str, str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    except Exception:
        return []
    text = re.sub(r"<[^>]+>", " ", xml)
    text = re.sub(r"\s+", " ", text).strip()
    out = []
    for m in LINE_RE.finditer(text):
        name = m.group("name").strip()
        qty = m.group("qty").strip()
        price = m.group("price").strip()
        if _is_valid_position(name, qty, price):
            out.append({"item": name, "quantity": qty, "price_for_one": price, "source": source + ":regex"})
    return out


def parse_document_url(url: str, title: str = "", referer: str = "") -> List[Dict[str, str]]:
    try:
        raw = fetch_bytes(url, timeout=60, referer=referer or url)
    except Exception:
        return []
    source = f"doc:{title[:45]}" if title else "doc"
    low = (url + " " + title).lower()

    if raw[:2] == b"PK":
        if low.endswith(".xlsx") or "xlsx" in title.lower():
            rows = parse_xlsx_tables(raw, source + ":xlsx")
            if rows:
                return rows
        rows = parse_docx_tables(raw, source + ":table")
        if rows:
            return rows
        return _regex_fallback(raw, source)

    if is_ole_doc(raw, title, url):
        docx_raw = convert_doc_to_docx(raw, title or "document.doc")
        if docx_raw:
            rows = parse_docx_tables(docx_raw, source + ":doc-table")
            if rows:
                return rows
            return _regex_fallback(docx_raw, source + ":doc")

    return []


def score_document_rows(rows: List[Dict[str, str]]) -> int:
    if not rows:
        return 0
    priced = sum(1 for r in rows if (_parse_price(r.get("price_for_one", "")) or 0) >= 100)
    qty = sum(1 for r in rows if (r.get("quantity") or "").strip())
    long_names = sum(1 for r in rows if len((r.get("item") or "")) > 180)
    short = sum(1 for r in rows if len((r.get("item") or "")) < 12)
    return priced * 25 + qty * 4 + len(rows) - long_names * 40 - short * 10


def parse_documents(doc_links: List[Dict], docs_page: str = "", limit: int = 6) -> List[Dict[str, str]]:
    best_rows: List[Dict[str, str]] = []
    best_score = 0
    for d in pick_documents(doc_links, limit=limit):
        title = d.get("title", "")
        rows = parse_document_url(d["url"], title, referer=docs_page)
        if not rows:
            continue
        score = score_document_rows(rows)
        if score > best_score:
            best_score = score
            best_rows = rows
    return best_rows
